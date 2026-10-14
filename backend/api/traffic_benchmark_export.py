from io import BytesIO
from typing import List

from fastapi import APIRouter
from fastapi.responses import Response
from pydantic import BaseModel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

router = APIRouter(prefix="/analytics", tags=["traffic-benchmark"])


class BenchmarkRow(BaseModel):
    method: str
    cars_before: int
    cars_rerouted: int
    cars_after: int
    congestion_after: float
    congestion_reduction_percent: float
    score: float
    prediction_accuracy_percent: float
    preferred: str


class BenchmarkExport(BaseModel):
    generated_at: str
    hotspot: str
    cars_before: int
    congestion_before: float
    preferred_method: str
    rows: List[BenchmarkRow]


@router.post("/traffic-benchmark-xlsx")
def export_traffic_benchmark(payload: BenchmarkExport):
    wb = Workbook()
    ws = wb.active
    ws.title = "Traffic Benchmark"

    ws["A1"] = "Quantum City - Traffic Optimization Benchmark"
    ws["A1"].font = Font(bold=True, size=16)

    ws["A3"] = "Generated At"
    ws["B3"] = payload.generated_at
    ws["A4"] = "Congested Road"
    ws["B4"] = payload.hotspot
    ws["A5"] = "Cars Before"
    ws["B5"] = payload.cars_before
    ws["A6"] = "Congestion Before"
    ws["B6"] = payload.congestion_before
    ws["A7"] = "Preferred Method"
    ws["B7"] = payload.preferred_method

    headers = [
        "Method",
        "Cars Before",
        "Cars Rerouted",
        "Cars After",
        "Congestion After",
        "Congestion Reduction %",
        "Score",
        "Prediction Accuracy %",
        "Preferred",
    ]

    start_row = 9

    for col, header in enumerate(headers, 1):
        cell = ws.cell(start_row, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163A5F")
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(payload.rows, start_row + 1):
        values = [
            row.method,
            row.cars_before,
            row.cars_rerouted,
            row.cars_after,
            row.congestion_after,
            row.congestion_reduction_percent,
            row.score,
            row.prediction_accuracy_percent,
            row.preferred,
        ]

        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)

    for row in range(start_row + 1, start_row + 1 + len(payload.rows)):
        ws.cell(row, 5).number_format = "0.0%"
        ws.cell(row, 6).number_format = "0.0%"
        ws.cell(row, 7).number_format = "0.0000"
        ws.cell(row, 8).number_format = "0.0%"

    widths = [22, 14, 15, 13, 19, 24, 12, 24, 12]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    ws.freeze_panes = "A10"

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="QuantumCity_Traffic_Benchmark.xlsx"'
        },
    )
